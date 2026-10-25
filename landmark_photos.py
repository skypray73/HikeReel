# -*- coding: utf-8 -*-
"""
HikeReel 地標照片（2D / 3D 共用）  —  MIT License

照片來源（擇一）：
  * 地標庫（.xlsx / .csv）加一欄「照片」，填照片路徑（相對路徑以地標庫所在資料夾為準）
  * 照片資料夾（預設為地標庫旁的 landmark_photos），檔名 = 地標名稱（中文或外文那一行皆可）
走到有照片的地標時停下，照片卡片淡入、緩慢推近；沒有照片的地標不顯示卡片、不停留。
"""
import os
import shutil

import numpy as np
from PIL import Image, ImageDraw

DEFAULT_FOLDER = "landmark_photos"


def default_folder(landmark_db):
    return os.path.join(os.path.dirname(os.path.abspath(landmark_db)), DEFAULT_FOLDER)


def save_photo(src, name, folder):
    """把使用者選的照片複製到照片資料夾，檔名用地標名稱的第一行；同名舊照片（其他副檔名）一併移除。"""
    os.makedirs(folder, exist_ok=True)
    stem = safe_stem(name)
    remove_photo(name, folder)
    dst = os.path.join(folder, stem + os.path.splitext(src)[1].lower())
    shutil.copy2(src, dst)
    return dst


def remove_photo(name, folder):
    stem = safe_stem(name).lower()
    if not os.path.isdir(folder):
        return
    for f in os.listdir(folder):
        st, ext = os.path.splitext(f)
        if ext.lower() in PHOTO_EXT and st.lower() == stem:
            os.remove(os.path.join(folder, f))


def safe_stem(name):
    first = str(name).split("\n")[0].strip()
    return "".join("_" if c in '<>:"/\\|?*' else c for c in first) or "landmark"


PHOTO_EXT = (".jpg", ".jpeg", ".png", ".webp", ".heic", ".heif", ".bmp")


def landmark_photo_table(path):
    """地標庫（.xlsx / .csv）裡若有「照片」欄，回傳 {名稱: 照片路徑}（相對路徑以地標庫所在資料夾為準）。"""
    out = {}
    try:
        rows = []
        if path.lower().endswith((".xlsx", ".xls")):
            import openpyxl
            ws = openpyxl.load_workbook(path, data_only=True).active
            allr = list(ws.iter_rows(values_only=True))
            head = [str(h).strip().lower() if h else "" for h in allr[0]]
            rows = [dict(zip(head, r)) for r in allr[1:]]
        elif path.lower().endswith(".csv"):
            import csv
            with open(path, newline="", encoding="utf-8-sig") as f:
                rows = [{k.strip().lower(): v for k, v in r.items()} for r in csv.DictReader(f)]
        base = os.path.dirname(os.path.abspath(path))
        for r in rows:
            name = next((r[k] for k in ("名稱", "name", "地標", "landmark") if r.get(k)), None)
            ph = next((r[k] for k in ("照片", "photo", "圖片", "image", "picture") if r.get(k)), None)
            if name and ph:
                ph = str(ph).strip()
                out[str(name)] = ph if os.path.isabs(ph) else os.path.join(base, ph)
    except Exception:
        pass
    return out


def find_photo(name, table, folder):
    """依序找：地標庫的「照片」欄 → 照片資料夾裡檔名等於地標名稱（任一行，不分大小寫）的圖片。"""
    if name in table and os.path.isfile(table[name]):
        return table[name]
    if not folder or not os.path.isdir(folder):
        return None
    keys = {name.replace("\n", " ").strip().lower(), name.replace("\n", "").strip().lower()}
    keys |= {t.strip().lower() for t in name.split("\n") if t.strip()}
    for f in sorted(os.listdir(folder)):
        stem, ext = os.path.splitext(f)
        if ext.lower() in PHOTO_EXT and stem.strip().lower() in keys:
            return os.path.join(folder, f)
    return None


def open_photo(path):
    from PIL import ImageOps
    if path.lower().endswith((".heic", ".heif")):
        try:
            import pillow_heif
            pillow_heif.register_heif_opener()
        except Exception:
            pass
    im = Image.open(path)
    return ImageOps.exif_transpose(im).convert("RGB")             # 依 EXIF 轉正（手機直拍照片）


class PhotoCard:
    """走到地標時顯示的照片卡片：自動依卡片比例置中裁切、淡入淡出、緩慢推近（Ken Burns）。"""
    ZOOM = 0.08

    def __init__(self, path, name, W, H, scale, font):
        from PIL import ImageOps
        land = W > H
        pad = int(min(W, H) * 0.035)
        self.cw = int(W * 0.40) if land else int(W * 0.86)
        self.ch = int(self.cw * 3 / 4)                               # 4:3 卡片
        self.x = W - pad - self.cw if land else (W - self.cw) // 2
        self.y = int(H * 0.12) if land else int(H * 0.13)
        big = (int(self.cw * (1 + self.ZOOM)), int(self.ch * (1 + self.ZOOM)))
        self.src = ImageOps.fit(open_photo(path), big, Image.LANCZOS, centering=(0.5, 0.45))   # 略偏上：保留天空與山稜
        self.lines = [line for line in str(name or "").splitlines() if line.strip()]
        self.scale = scale
        self.font = font

    def draw(self, frame, u, alpha):
        """u：停留進度 0~1；alpha：淡入淡出 0~1。"""
        if alpha <= 0.01:
            return frame
        z = 1 + self.ZOOM * u
        bw, bh = self.src.size
        cw, ch = int(bw / z), int(bh / z)
        x0, y0 = (bw - cw) // 2, int((bh - ch) * 0.45)
        photo = self.src.crop((x0, y0, x0 + cw, y0 + ch)).resize((self.cw, self.ch), Image.LANCZOS)
        fnt = self.font
        cap_h = int(fnt.size * 1.35 * len(self.lines) + 16 * self.scale) if self.lines else 0
        bd = max(3, int(4 * self.scale))
        card = Image.new("RGBA", (self.cw + 2 * bd, self.ch + cap_h + 2 * bd), (255, 255, 255, 255))
        card.paste(photo, (bd, bd))
        dr = ImageDraw.Draw(card)
        if self.lines:
            dr.rectangle([bd, self.ch + bd, self.cw + bd, self.ch + bd + cap_h], fill=(15, 20, 40, 255))
        for k, t in enumerate(self.lines):
            dr.text((card.width / 2, self.ch + bd + 8 * self.scale + fnt.size * 1.35 * k), t, font=fnt,
                    fill=(255, 255, 255) if k == 0 else (200, 205, 225), anchor="ma")
        mask = Image.new("L", card.size, 0)
        ImageDraw.Draw(mask).rounded_rectangle([0, 0, card.width - 1, card.height - 1], radius=int(14 * self.scale),
                                               fill=int(255 * alpha))
        im = Image.fromarray(frame).convert("RGBA")
        sh = Image.new("RGBA", im.size, (0, 0, 0, 0))               # 陰影
        ImageDraw.Draw(sh).rounded_rectangle([self.x + 6, self.y + 8, self.x + card.width + 6, self.y + card.height + 8],
                                             radius=int(14 * self.scale), fill=(0, 0, 0, int(90 * alpha)))
        im = Image.alpha_composite(im, sh)
        im.paste(card, (self.x, self.y), mask)
        return np.asarray(im.convert("RGB"))


def build_schedule(L, frames_move, holds):
    """把「沿路線移動」的幀和「在地標停留」的幀排在一起。
    holds：[(距離 m, 停留幀數, 物件)]；回傳 (每幀的路線距離, [(開始幀, 結束幀, 物件)])。"""
    s_move = np.linspace(0, L, frames_move)
    out, ranges, hi = [], [], 0
    holds = sorted(holds, key=lambda h: h[0])
    for v in s_move:
        while hi < len(holds) and holds[hi][0] <= v:
            d, n, obj = holds[hi]
            ranges.append((len(out), len(out) + n, obj))
            out += [d] * n
            hi += 1
        out.append(v)
    return np.array(out), ranges


def list_landmark_names(path):
    """讀出地標庫裡所有地標名稱（給「管理地標照片」視窗用）。"""
    names = []
    try:
        if path.lower().endswith((".xlsx", ".xls")):
            import openpyxl
            ws = openpyxl.load_workbook(path, data_only=True).active
            rows = list(ws.iter_rows(values_only=True))
            head = [str(h).strip().lower() if h else "" for h in rows[0]]
            col = next((i for i, h in enumerate(head) if h in ("名稱", "name", "地標", "landmark")), None)
            lat = next((i for i, h in enumerate(head) if h in ("緯度", "lat", "latitude", "座標", "coord", "coordinates", "gps", "經緯度")), None)
            for r in rows[1:]:
                if col is None or r[col] in (None, ""):
                    continue
                if lat is not None and r[lat] is not None:
                    try:                                  # 跳過說明列（座標欄不是數字）
                        float(str(r[lat]).replace(",", " ").split()[0])
                    except (ValueError, IndexError):
                        continue
                names.append(str(r[col]))
        elif path.lower().endswith(".csv"):
            import csv
            with open(path, newline="", encoding="utf-8-sig") as f:
                for r in csv.DictReader(f):
                    r = {k.strip().lower(): v for k, v in r.items()}
                    n = next((r[k] for k in ("名稱", "name", "地標", "landmark") if r.get(k)), None)
                    if n:
                        names.append(n)
    except Exception:
        pass
    seen, out = set(), []
    for n in names:
        if n not in seen:
            seen.add(n)
            out.append(n)
    return out
